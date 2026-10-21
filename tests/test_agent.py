from decimal import Decimal
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app.harness import ExtractionHarness, HarnessCase
from app.parsers.ocr_client import OCRClient
from app.parsers.router import DocumentRouter
from app.schemas import AuditTrail, BillRecord, ChargeDetails, DocumentInfo, FinancialAttributes, TransactionParty
from app.services.agent import ShippingBillAgent
from app.services.matcher import BillMatcher


def build(amount: str) -> BillRecord:
    return BillRecord(
        operation="CREATE",
        document_info=DocumentInfo(reference_no="DN-1"),
        transaction_party=TransactionParty(issuer="Destination", recipient="Origin"),
        charge_details=ChargeDetails(raw_charge_item="Handling Fee", charge_item="Handling Charge", amount=Decimal(amount), currency="USD", payment_term="PP"),
        financial_attributes=FinancialAttributes(accounting_entity="Origin", posting_type="AP"),
        audit_trail=AuditTrail(),
    )


def test_matcher_prefers_reference_and_charge():
    matches = BillMatcher().find(build("40"), [build("30")])
    assert len(matches) == 1
    assert matches[0].rule == "reference_no + charge_item"


def test_rule_extractor_creates_one_record_per_charge():
    agent = ShippingBillAgent()
    extracted = agent.extractor.extract("Debit Note DN-202609-001 2026-09-20\nHandling Fee USD 40.00 PP\nDoc Fee USD 20.00 CC", "test")
    records = agent._to_records([(extracted, ["test"])], "Origin Agent")
    assert len(records) == 2
    assert records[0].charge_details.charge_item == "Handling Charge"


def test_rule_extractor_reads_pdf_table_rows_with_header_currency():
    agent = ShippingBillAgent()
    text = "DEBIT NOTE\nD/C No. : DC-108478\nD/C DATE 10-24-2026\nCURRENCY USD\nHANDLING CHARGE - 65.00 1.00 C 65.00\n4 ADD'L DETENTION FREE TIME CHARGE - 1,310.40 1.00 C 1,310.40"
    extracted = agent.extractor.extract(text, "test.pdf")
    assert extracted.document_info.reference_no == "DC-108478"
    assert len(extracted.charges) == 2
    assert extracted.charges[1].amount == Decimal("1310.40")


def test_update_inherits_unmodified_fields():
    agent = ShippingBillAgent()
    previous = build("30")
    previous.charge_details.payment_term = "PP"
    incoming = build("40")
    incoming.charge_details.payment_term = None
    agent._decide(incoming, [previous])
    assert incoming.operation == "UPDATE"
    assert incoming.charge_details.payment_term == "PP"


def test_ar_ap_needs_entity_party_and_direction():
    agent = ShippingBillAgent()
    assert agent._posting_type("Origin", "Destination", "Origin", "RECEIVABLE") == "AR"
    assert agent._posting_type("Origin", "Destination", "Destination", "PAYABLE") == "AP"
    assert agent._posting_type("Origin", "Destination", "Origin", "PAYABLE") is None


def test_declared_total_conflict_requires_review():
    agent = ShippingBillAgent()
    result = agent.extractor.extract("D/C No. : DC-100\nCURRENCY USD\nTOTAL AMOUNT 100.00\nHANDLING CHARGE - 50.00 1.00 C 50.00", "test")
    records = agent._to_records([(result, ["test"])], "Origin")
    assert records[0].operation == "PENDING_REVIEW"
    assert "conflicts with declared" in records[0].audit_trail.validations[0]


def test_excel_and_docx_tables_are_parsed():
    workbook = Workbook()
    workbook.active.append(["Fee", "Amount"])
    workbook.active.append(["Handling Fee", 40])
    excel_bytes = BytesIO()
    workbook.save(excel_bytes)
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Fee"
    table.rows[0].cells[1].text = "Amount"
    word_bytes = BytesIO()
    document.save(word_bytes)
    router = DocumentRouter()
    assert "Handling Fee" in router._parse_excel_stream("test.xlsx", excel_bytes.getvalue()).text
    assert "Amount" in router._parse_docx_stream("test.docx", word_bytes.getvalue()).text


def test_harness_keeps_pdf_table_regression_case_green():
    case = HarnessCase(
        name="pdf-table",
        text="D/C No. : DC-1\nCURRENCY USD\nHANDLING CHARGE - 40.00 1.00 C 40.00",
        expected_charge_count=1,
        expected_amounts=(Decimal("40.00"),),
    )
    result = ExtractionHarness().run(case)
    assert result.passed, result.details


def test_eml_image_attachment_uses_ocr_url_mapping():
    class FakeOCR:
        def extract(self, url):
            assert url == "https://files.example.com/bill.png"
            return {"text": "Handling Fee USD 40.00"}

        @staticmethod
        def to_text(payload):
            return payload["text"]

    path = Path(__file__).parent / "fixtures" / "image_attachment.eml"
    parsed = DocumentRouter(ocr_client=FakeOCR()).parse(path, attachment_urls={"bill.png": "https://files.example.com/bill.png"})
    assert parsed.attachments[0].kind == "image"
    assert "Handling Fee" in parsed.attachments[0].text


def test_legacy_doc_uses_docx_conversion_when_available(monkeypatch):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Handling Fee"
    table.rows[0].cells[1].text = "40"
    converted = BytesIO()
    document.save(converted)
    monkeypatch.setattr("app.parsers.router.LegacyDocConverter.to_docx", lambda _: converted.getvalue())
    parsed = DocumentRouter._parse_doc_stream("legacy.doc", b"binary-doc")
    assert parsed.kind == "word"
    assert parsed.tables[0][0] == ["Handling Fee", "40"]


def test_ocr_client_uses_required_api_contract(monkeypatch):
    captured = {}

    class Response:
        def raise_for_status(self):
            return None

        @staticmethod
        def json():
            return {"text": "Handling Fee USD 40.00"}

    def fake_post(url, json, timeout):
        captured.update(url=url, json=json, timeout=timeout)
        return Response()

    monkeypatch.setattr("app.parsers.ocr_client.requests.post", fake_post)
    result = OCRClient("http://ocr.example/process").extract("https://files.example.com/bill.png")
    assert result["text"] == "Handling Fee USD 40.00"
    assert captured["url"] == "http://ocr.example/process"
    assert captured["json"] == {"file": "https://files.example.com/bill.png", "task": "orderChange", "callBack": "", "taskType": ""}
